#!/usr/bin/env python3
"""Build candidates.xlsx from the per-candidate .json sidecars written by the Kiro hiring-agent.

Usage: python resumes_to_excel.py [resumes_dir]
Re-runnable: the workbook is rebuilt from whatever is on disk.
"""
import json
import subprocess
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "openpyxl"])
    import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(sys.argv[1] if len(sys.argv) > 1 else "/mnt/c/Users/apoor/Documents/resumes")
COLS = ["job_title", "candidate_name", "location", "phone", "email", "current_title",
        "current_company", "total_experience", "education", "key_skills", "expected_salary",
        "applied_date", "indeed_profile_url", "notes"]


def win_link(p: Path) -> str:
    s = str(p)
    if s.startswith("/mnt/") and len(s) > 6:
        s = s[5].upper() + ":" + s[6:]
    return "file:///" + s.replace("\\", "/")


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidates"
    ws.append(["Job folder", "Folder link", *[c.replace("_", " ").title() for c in COLS], "Resume file", "Resume link"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")

    n = 0
    for folder in sorted(p for p in ROOT.iterdir() if p.is_dir() and not p.name.startswith("_")):
        for js in sorted(folder.glob("*.json")):
            try:
                d = json.loads(js.read_text(encoding="utf-8"))
            except json.JSONDecodeError:
                print("skip unreadable", js)
                continue
            resume = next((f for f in folder.iterdir() if f.stem == js.stem and f.suffix.lower() != ".json"), None)
            ws.append([folder.name, "Open folder", *[d.get(c, "") for c in COLS],
                       resume.name if resume else "MISSING", "Open resume" if resume else ""])
            r = ws.max_row
            ws.cell(r, 2).hyperlink = win_link(folder)
            ws.cell(r, 2).style = "Hyperlink"
            if resume:
                ws.cell(r, ws.max_column).hyperlink = win_link(resume)
                ws.cell(r, ws.max_column).style = "Hyperlink"
            n += 1

    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(12, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    out = ROOT / "candidates.xlsx"
    wb.save(out)
    print(f"Wrote {out} ({n} candidates)")


if __name__ == "__main__":
    main()
